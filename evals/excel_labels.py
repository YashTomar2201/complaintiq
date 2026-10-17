"""Excel version of the labeling page: make a workbook to fill in, then import it.

  python -m evals.excel_labels make     # writes data/gold_labeling.xlsx (100 complaints, dropdowns, rubric sheet)
  python -m evals.excel_labels import   # reads your filled workbook into evals/gold_labels.csv (checks every row)

You fill the three yellow columns (root cause, severity, vulnerable); notes are optional.
Importing is safe to repeat: it adds or replaces labels by complaint id.
"""
import json
import math
import sys

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from evals.gold import GOLD_SET, GOLD_TEXTS, load_labels, save_label
from src.ai.schema import ROOT_CAUSE_HELP, SEVERITY_HELP, VULNERABLE_HELP
from src.ingest.cfpb import ROOT

XLSX = ROOT / "data" / "gold_labeling.xlsx"
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")   # cells you fill in
HEADER = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TEXT_WIDTH = 100          # characters per line in the complaint-text column
MAX_ROW_HEIGHT = 400      # Excel's limit is 409 points; very long texts can be read by clicking the cell


def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def make_workbook(path=XLSX):
    gold = pd.read_csv(GOLD_SET, dtype=str)
    texts = json.loads(GOLD_TEXTS.read_text(encoding="utf-8"))
    wb = Workbook()

    # ---------- sheet 1: read me ----------
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    rows = [
        ("How to fill this workbook", None),
        ("1", "Go to the 'Labels' sheet. Read each complaint (column E) and fill the three YELLOW columns: F, G, H."),
        ("2", "Use the dropdown arrow in each yellow cell. Typing anything else is rejected."),
        ("3", "Use the 'Rubric' sheet for the definitions. They are the SAME definitions the AI was given."),
        ("4", "Label from your own reading. Do NOT ask an AI: the test only means something if the answers are yours."),
        ("5", "You can save and continue any time. Leave a row empty to skip it for now."),
        ("6", "When finished, save and close the file, then tell Claude. Do not rename the file or change columns A-E."),
        (None, None),
        ("Example (not in your data)", None),
        ("Complaint text", "I am 78 and someone used my card for {$2,400}. Capital One froze my account and I cannot pay rent."),
        ("Root cause (F)", "unauthorized_transaction"),
        ("Severity (G)", "5   (money lost, vulnerable consumer, severe hardship)"),
        ("Vulnerable? (H)", "Yes   (elderly, cannot pay rent)"),
        (None, None),
        ("Notes", "Text like XXXX or <PERSON> is privacy masking. Dollar amounts appear as {$123.45}. "
                  "Very long complaints may be cut off in the row: click the cell to read all of it in the formula bar."),
    ]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = f(bold=True, size=12 if i == 1 else 10)
        if b:
            c = ws.cell(i, 2, b)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in (9,):
        ws.cell(r, 1).font = f(bold=True, size=11)
    for r in (11, 12, 13):
        ws.cell(r, 2).fill = YELLOW

    # ---------- sheet 2: labels ----------
    ls = wb.create_sheet("Labels")
    heads = ["#", "Complaint ID", "Product", "Issue the consumer chose", "Complaint text (masked)",
             "Root cause (YOU)", "Severity 1-5 (YOU)", "Vulnerable? (YOU)", "Notes (optional)"]
    widths = [5, 13, 24, 30, TEXT_WIDTH, 24, 12, 14, 30]
    for col, (h, w) in enumerate(zip(heads, widths), 1):
        c = ls.cell(1, col, h)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEADER if col < 6 else PatternFill("solid", fgColor="BF8F00")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ls.column_dimensions[ls.cell(1, col).column_letter].width = w
    ls.row_dimensions[1].height = 32
    for i, g in enumerate(gold.itertuples(), 2):
        t = texts[g.complaint_id]
        vals = [int(g.position), g.complaint_id, t["product"], t["issue"], t["text"]]
        for col, v in enumerate(vals, 1):
            c = ls.cell(i, col, v)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        ls.cell(i, 2).number_format = "@"  # complaint ids stay text
        for col in range(6, 10):
            c = ls.cell(i, col)
            c.fill, c.font, c.border = YELLOW, f(), BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        lines = sum(max(1, math.ceil(len(p) / (TEXT_WIDTH * 1.05))) for p in t["text"].split("\n"))
        ls.row_dimensions[i].height = min(max(30, 13.5 * lines + 6), MAX_ROW_HEIGHT)
    last = len(gold) + 1
    ls.freeze_panes = "F2"

    # ---------- sheet 3: rubric (also the source of the dropdown list) ----------
    rs = wb.create_sheet("Rubric")
    rs.column_dimensions["A"].width = 26
    rs.column_dimensions["B"].width = 120
    rs["A1"] = "Root cause: pick ONE (what went wrong)"
    rs["A1"].font = f(bold=True, size=12)
    for i, (k, v) in enumerate(ROOT_CAUSE_HELP.items(), 2):
        rs.cell(i, 1, k).font = f(bold=True)
        rs.cell(i, 2, v).font = f()
    n_rc = len(ROOT_CAUSE_HELP)
    r0 = n_rc + 4
    rs.cell(r0, 1, "Severity: pick the HIGHEST level that applies").font = f(bold=True, size=12)
    for i, (k, v) in enumerate(SEVERITY_HELP.items(), r0 + 1):
        rs.cell(i, 1, k).font = f(bold=True)
        rs.cell(i, 1).alignment = Alignment(horizontal="left")
        rs.cell(i, 2, v).font = f()
        rs.cell(i, 2).alignment = Alignment(wrap_text=True)
    r1 = r0 + len(SEVERITY_HELP) + 3
    rs.cell(r1, 1, "Vulnerable consumer?").font = f(bold=True, size=12)
    rs.cell(r1 + 1, 1, "Yes").font = f(bold=True)
    rs.cell(r1 + 1, 2, VULNERABLE_HELP).font = f()
    rs.cell(r1 + 1, 2).alignment = Alignment(wrap_text=True)
    rs.cell(r1 + 2, 1, "No").font = f(bold=True)
    rs.cell(r1 + 2, 2, "none of the above is stated or clearly shown").font = f()

    # ---------- dropdowns ----------
    dv_root = DataValidation(type="list", formula1=f"=Rubric!$A$2:$A${n_rc + 1}", allow_blank=True,
                             showErrorMessage=True, errorTitle="Pick from the list", error="Choose one of the root causes in the dropdown.")
    dv_sev = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Severity must be 1, 2, 3, 4 or 5.")
    dv_vul = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Answer Yes or No.")
    for dv, col in ((dv_root, "F"), (dv_sev, "G"), (dv_vul, "H")):
        ls.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    wb.active = 1  # open on the Labels sheet
    path.parent.mkdir(exist_ok=True)
    wb.save(path)
    return path


# ---------- import ----------
def read_filled(path=XLSX):
    """Returns (valid label dicts, list of problems, number of rows left empty)."""
    df = pd.read_excel(path, sheet_name="Labels", dtype=str).fillna("")
    df.columns = ["pos", "complaint_id", "product", "issue", "text", "root_cause", "severity", "vulnerable", "notes"]
    known = set(pd.read_csv(GOLD_SET, dtype=str).complaint_id)
    good, problems, empty = [], [], 0
    for r in df.itertuples():
        cid = str(r.complaint_id).strip()
        rc, sev, vul = r.root_cause.strip(), r.severity.strip(), r.vulnerable.strip()
        if cid not in known:
            problems.append(f"row {r.Index + 2}: complaint id {cid!r} is not in the gold set (do not edit columns A-E)")
        elif not (rc or sev or vul):
            empty += 1
        elif not (rc and sev and vul):
            problems.append(f"row {r.Index + 2} (#{r.pos}): fill all three yellow cells or leave all empty")
        elif rc not in ROOT_CAUSE_HELP:
            problems.append(f"row {r.Index + 2} (#{r.pos}): root cause {rc!r} is not one of the allowed values")
        elif sev.split(".")[0] not in {"1", "2", "3", "4", "5"}:
            problems.append(f"row {r.Index + 2} (#{r.pos}): severity {sev!r} must be 1-5")
        elif vul.lower() not in ("yes", "no"):
            problems.append(f"row {r.Index + 2} (#{r.pos}): vulnerable {vul!r} must be Yes or No")
        else:
            good.append(dict(complaint_id=cid, root_cause=rc, severity=int(float(sev)), vulnerable=vul.lower() == "yes",
                             notes=r.notes.strip()))
    return good, problems, empty


def import_labels(path=XLSX, labels_path=None):
    kw = {"path": labels_path} if labels_path else {}
    good, problems, empty = read_filled(path)
    before = len(load_labels(**kw))
    for g in good:
        save_label(g["complaint_id"], g["root_cause"], g["severity"], g["vulnerable"], g["notes"], **kw)
    return dict(imported=len(good), empty=empty, problems=problems, total_labels=len(load_labels(**kw)), before=before)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "make":
        if XLSX.exists():
            raise SystemExit(f"{XLSX} already exists (it may hold your answers). Delete it first if you really want a fresh copy.")
        print("created", make_workbook())
    elif cmd == "import":
        res = import_labels()
        print(f"imported {res['imported']} labels ({res['empty']} rows still empty); labels file now has {res['total_labels']}")
        for p in res["problems"]:
            print("PROBLEM:", p)
    else:
        raise SystemExit(__doc__)
