"""Hand-check of the regulation suggestions: 20 complaints, a workbook for the human to judge them.

  python -m evals.regulation_check make     # runs the suggester on 20 gold complaints, writes data/regulation_check.xlsx
  python -m evals.regulation_check import   # reads the judged workbook, prints the percentages, saves evals/results/regulation_check.json

The 20 complaints come from the gold set (the ones you already know), spread over root causes (your own labels).
"""
import json
import sys

import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from evals.gold import GOLD_TEXTS, load_labels
from src.ingest.cfpb import DB_PATH, ROOT

XLSX = ROOT / "data" / "regulation_check.xlsx"
RESULT = ROOT / "evals" / "results" / "regulation_check.json"
N, SEED = 20, 11
VERDICTS = ["Sensible", "Partly sensible", "Not sensible", "None suggested - correct", "None suggested - a rule was missed"]
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def pick_complaints(labels, n=N, seed=SEED):
    """Round-robin over the human root-cause labels so every kind of complaint is represented."""
    df = pd.DataFrame(labels.values()).sample(frac=1, random_state=seed)
    groups = {rc: list(g.complaint_id) for rc, g in df.groupby("root_cause")}
    chosen = []
    while len(chosen) < n and any(groups.values()):
        for rc in sorted(groups):
            if groups[rc] and len(chosen) < n:
                chosen.append(groups[rc].pop())
    return chosen


def make(version="v2"):
    from src.ai.llm import get_aux_client as get_client
    from src.ai.regulations_rag import suggest
    from src.ai.embeddings import collection
    labels = load_labels()
    texts = json.loads(GOLD_TEXTS.read_text(encoding="utf-8"))
    con = duckdb.connect(str(DB_PATH), read_only=True)
    triage = con.execute("""SELECT complaint_id, summary, root_cause FROM complaint_triage
                            WHERE prompt_version = $v AND status = 'ok'""", {"v": version}).df().set_index("complaint_id")
    con.close()
    client, coll = get_client(), collection("regulations")
    rows = []
    for cid in pick_complaints(labels):
        t = triage.loc[cid]
        out = suggest(client, t.summary, t.root_cause, texts[cid]["product"], coll)
        rows.append(dict(cid=cid, product=texts[cid]["product"], text=texts[cid]["text"], summary=t.summary,
                         root_cause=t.root_cause, suggestions=out["suggestions"], rejected=out["rejected"]))
    return write_workbook(rows)


def fmt(suggestions):
    if not suggestions:
        return "(none suggested)"
    return "\n\n".join(f"{s['citation']}: {s['title']}\n{s['explanation']}" for s in suggestions)


def write_workbook(rows, path=XLSX):
    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    heads = ["#", "Complaint ID", "Complaint text (masked)", "AI summary / root cause", "Suggested regulation sections",
             "Your verdict (YOU)", "Notes (optional)"]
    widths = [5, 12, 80, 40, 70, 28, 30]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER if c < 6 else PatternFill("solid", fgColor="BF8F00")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w
    for i, r in enumerate(rows, 2):
        vals = [i - 1, r["cid"], r["text"][:1500], f"{r['summary']}\n[{r['root_cause']}]", fmt(r["suggestions"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.font, cell.border = Font(name=FONT, size=10), BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 2).number_format = "@"
        for c in (6, 7):
            cell = ws.cell(i, c)
            cell.fill, cell.border, cell.font = YELLOW, BOX, Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        n_lines = max(len(r["text"][:1500]) // 85, fmt(r["suggestions"]).count("\n") + len(fmt(r["suggestions"])) // 75, 4)
        ws.row_dimensions[i].height = min(15 * n_lines + 8, 400)
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick from the list", error="Choose one of the verdicts in the dropdown.")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(rows) + 1}")
    ws.freeze_panes = "C2"
    rm = wb.create_sheet("Read me", 0)
    rm.column_dimensions["A"].width = 120
    for i, line in enumerate([
        "How to judge the regulation suggestions",
        "For each complaint, read the text and the suggested sections, then choose a verdict in the yellow column F.",
        "Sensible = the suggested section really may apply.   Partly sensible = some do, some do not, or it is only loosely related.",
        "Not sensible = the suggestion does not fit the complaint.",
        "None suggested - correct = no section was suggested and none applies.   None suggested - a rule was missed = no section was shown but one applies.",
        "You are the expert here: judge from your own reading. This is a suggestion for compliance review, not legal advice.",
        "Save and close the file when you are done, then tell Claude.",
    ], 1):
        rm.cell(i, 1, line).font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
        rm.cell(i, 1).alignment = Alignment(wrap_text=True)
    wb.active = 1
    path.parent.mkdir(exist_ok=True)
    wb.save(path)
    return path


def score(path=XLSX):
    df = pd.read_excel(path, sheet_name="Check", dtype=str).fillna("")
    v = df.iloc[:, 5].str.strip()
    done = v[v != ""]
    bad = sorted(set(done) - set(VERDICTS))
    if bad:
        raise SystemExit(f"Unknown verdicts: {bad}")
    n_sug = int(done.isin(["Sensible", "Partly sensible", "Not sensible"]).sum())
    res = dict(judged=len(done), of=len(df), counts=done.value_counts().to_dict(),
               sensible_pct_of_suggestions=round(100 * (done == "Sensible").sum() / n_sug, 1) if n_sug else None,
               sensible_or_partly_pct=round(100 * done.isin(["Sensible", "Partly sensible"]).sum() / n_sug, 1) if n_sug else None,
               no_suggestion=int(done.str.startswith("None").sum()), missed=int((done == "None suggested - a rule was missed").sum()))
    RESULT.parent.mkdir(exist_ok=True)
    RESULT.write_text(json.dumps(res, indent=2), encoding="utf-8")
    return res


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "make":
        print("created", make())
    elif cmd == "import":
        print(json.dumps(score(), indent=2))
    else:
        raise SystemExit(__doc__)
