import pandas as pd
from openpyxl import load_workbook

from evals import regulation_check as rc


def test_pick_covers_root_causes_round_robin():
    labels = {str(i): dict(complaint_id=str(i), root_cause=c) for i, c in enumerate(["a"] * 10 + ["b"] * 10 + ["c"] * 2)}
    chosen = rc.pick_complaints(labels, n=6)
    causes = [labels[c]["root_cause"] for c in chosen]
    assert sorted(causes) == ["a", "a", "b", "b", "c", "c"] and len(set(chosen)) == 6


def test_workbook_roundtrip_and_scoring(tmp_path, monkeypatch):
    rows = [dict(cid=str(i), product="Credit card", text="t", summary="s", root_cause="billing_error",
                 suggestions=[dict(citation="12 CFR 1026.13", title="Billing error resolution", explanation="x. y.")] if i < 4 else [],
                 rejected=[]) for i in range(5)]
    path = rc.write_workbook(rows, tmp_path / "w.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Read me", "Check"] and "12 CFR 1026.13: Billing error resolution" in wb["Check"]["E2"].value
    assert wb["Check"]["E6"].value == "(none suggested)"
    ws = wb["Check"]
    for r, v in zip(range(2, 7), ["Sensible", "Sensible", "Partly sensible", "Not sensible", "None suggested - correct"]):
        ws.cell(r, 6, v)
    wb.save(path)
    monkeypatch.setattr(rc, "RESULT", tmp_path / "res.json")
    res = rc.score(path)
    assert res["judged"] == 5 and res["sensible_pct_of_suggestions"] == 50.0 and res["sensible_or_partly_pct"] == 75.0
    assert res["no_suggestion"] == 1 and res["missed"] == 0
