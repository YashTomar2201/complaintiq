import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from evals import excel_labels as xl
from evals.gold import load_labels


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    gold = tmp_path / "gold_set.csv"
    texts = tmp_path / "texts.json"
    pd.DataFrame({"position": [1, 2, 3], "complaint_id": ["111", "222", "333"], "product": ["Credit card"] * 3}).to_csv(gold, index=False)
    texts.write_text(json.dumps({c: {"product": "Credit card", "issue": "Fees", "text": f"text for {c} {{$5.00}}"} for c in ("111", "222", "333")}))
    monkeypatch.setattr(xl, "GOLD_SET", gold)
    monkeypatch.setattr(xl, "GOLD_TEXTS", texts)
    return xl.make_workbook(tmp_path / "w.xlsx")


def fill(path, rows):
    wb = load_workbook(path)
    ws = wb["Labels"]
    for r, (rc, sev, vul) in rows.items():
        ws.cell(r, 6, rc), ws.cell(r, 7, sev), ws.cell(r, 8, vul)
    wb.save(path)


def test_workbook_has_dropdowns_rubric_and_hides_nothing_secret(workbook):
    wb = load_workbook(workbook)
    assert wb.sheetnames == ["Read me", "Labels", "Rubric"]
    ls = wb["Labels"]
    assert ls["B2"].value == "111" and "text for 111" in ls["E2"].value
    assert len(ls.data_validations.dataValidation) == 3      # root cause, severity, vulnerable dropdowns
    assert ls["F1"].value.startswith("Root cause")


def test_import_checks_every_row(workbook, tmp_path):
    fill(workbook, {2: ("billing_error", 3, "No"), 3: ("made_up", 4, "No"), 4: ("fraud_scam", 5, None)})
    out = tmp_path / "labels.csv"
    res = xl.import_labels(workbook, labels_path=out)
    assert res["imported"] == 1 and len(res["problems"]) == 2
    assert any("not one of the allowed values" in p for p in res["problems"])
    assert any("fill all three" in p for p in res["problems"])
    assert load_labels(out)["111"]["root_cause"] == "billing_error"


def test_import_is_repeatable_and_replaces(workbook, tmp_path):
    out = tmp_path / "labels.csv"
    fill(workbook, {2: ("billing_error", 3, "No")})
    xl.import_labels(workbook, labels_path=out)
    fill(workbook, {2: ("fees_interest", 2, "Yes")})  # user changes their mind
    xl.import_labels(workbook, labels_path=out)
    labels = load_labels(out)
    assert len(labels) == 1 and labels["111"]["root_cause"] == "fees_interest" and labels["111"]["vulnerable_consumer"] == "True"
